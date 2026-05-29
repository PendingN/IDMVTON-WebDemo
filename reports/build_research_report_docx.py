from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "bao-cao-ung-dung-ai-thu-do-ao.md"
DOCX_PATH = ROOT / "bao-cao-ung-dung-ai-thu-do-ao.docx"


def set_font(run, size=13, bold=False, italic=False, color=None, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_style_font(style, size=13, bold=False, italic=False, color=None):
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    if color:
        font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def configure_styles(doc):
    styles = doc.styles
    set_style_font(styles["Normal"], 13)
    normal_pf = styles["Normal"].paragraph_format
    normal_pf.first_line_indent = Inches(0.35)
    normal_pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal_pf.line_spacing = 1.3
    normal_pf.space_after = Pt(6)

    for name, size, before, after in [
        ("Title", 16, 0, 12),
        ("Heading 1", 13, 14, 8),
        ("Heading 2", 13, 10, 6),
        ("Heading 3", 13, 8, 4),
    ]:
        set_style_font(styles[name], size=size, bold=True)
        pf = styles[name].paragraph_format
        pf.first_line_indent = None
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.2
        pf.keep_with_next = True

    set_style_font(styles["List Bullet"], 13)
    set_style_font(styles["List Number"], 13)


def configure_section(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for element in [
        ("begin", None),
        (None, "PAGE"),
        ("end", None),
    ]:
        if element[0]:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), element[0])
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = element[1]
            run._r.append(instr)
    set_font(run, 13)


def add_formatted_text(paragraph, text, base_size=13, bold=False, italic=False):
    # Supports simple **bold** and `code` inline markup.
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, base_size, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, base_size, bold=True, italic=italic)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, base_size, name="Courier New")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, base_size, bold=bold, italic=italic)


def add_para(doc, text, style=None, align=None, bold=False, italic=False, refs=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    add_formatted_text(p, text, bold=bold, italic=italic)
    if refs:
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
    return p


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc, table_lines):
    rows = parse_table(table_lines)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.first_line_indent = None
            para.paragraph_format.space_after = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(cell_text) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_text(para, cell_text, base_size=11, bold=(i == 0))
    doc.add_paragraph()


def build_docx():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), "Báo cáo nghiên cứu")

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    # Cover page.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TRƯỜNG/ĐƠN VỊ: ........................................\nKHOA/BỘ MÔN: ........................................")
    set_font(run, 13, bold=True)
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CÔNG TRÌNH NGHIÊN CỨU KHOA HỌC")
    set_font(run, 15, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title.upper())
    set_font(run, 16, bold=True)
    doc.add_paragraph()
    doc.add_paragraph()
    for label in [
        "Tác giả/nhóm tác giả: ........................................",
        "Lớp/Khoa: ........................................",
        "Giảng viên hướng dẫn: ........................................",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_font(run, 13)
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("...................., 2026")
    set_font(run, 13, bold=True)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    set_page_number_start(section, 1)
    add_page_number(section.header.paragraphs[0])

    body_started = False
    in_refs = False
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("## Tóm tắt công trình"):
            body_started = True
        if not body_started:
            continue

        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_buf.append(line)
            continue
        flush_table()

        if not line.strip():
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            in_refs = heading.startswith("6.1.")
            add_para(doc, heading.upper() if heading.startswith("Tóm tắt") else heading, style="Heading 1")
        elif line.startswith("### "):
            heading = line[4:].strip()
            in_refs = heading.startswith("6.1.")
            add_para(doc, heading, style="Heading 2")
        elif line.startswith("#### "):
            add_para(doc, line[5:].strip(), style="Heading 3")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = None
            add_formatted_text(p, line[2:].strip())
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = None
            add_formatted_text(p, re.sub(r"^\d+\.\s+", "", line).strip())
        elif line.startswith("Công thức ") or line.startswith("`") or line.startswith("Hình ") or line.startswith("Bảng "):
            is_caption = line.startswith("Hình ") or line.startswith("Bảng ")
            clean = line.strip("`")
            add_para(
                doc,
                clean,
                align=WD_ALIGN_PARAGRAPH.CENTER if is_caption else None,
                bold=is_caption,
                italic=is_caption,
            )
        else:
            add_para(doc, line, refs=in_refs)

    flush_table()
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
